"""Interchangeable readers; format-specific libraries load only when used."""
from abc import ABC, abstractmethod


class DocumentReader(ABC):
    @abstractmethod
    def read(self, content: bytes) -> str:
        """Extract plain text without retaining the uploaded bytes."""


class TextReader(DocumentReader):
    def read(self, content: bytes) -> str:
        from .txt_reader import read_txt
        return read_txt(content)


class PDFReader(DocumentReader):
    def read(self, content: bytes) -> str:
        from .pdf_reader import read_pdf
        return read_pdf(content)


class WordReader(DocumentReader):
    def read(self, content: bytes) -> str:
        from .docx_reader import read_docx
        return read_docx(content)


class PowerPointReader(DocumentReader):
    def read(self, content: bytes) -> str:
        from .pptx_reader import read_pptx
        return read_pptx(content)


class SpreadsheetReader(DocumentReader):
    def read(self, content: bytes) -> str:
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        blocks = []
        try:
            for sheet in workbook:
                blocks.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    text = "\t".join("" if value is None else str(value) for value in row)
                    if text.strip(): blocks.append(text)
        finally:
            workbook.close()
        return "\n".join(blocks)
