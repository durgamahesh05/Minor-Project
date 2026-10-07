from io import BytesIO

from pptx import Presentation


def read_pptx(content: bytes) -> str:
    presentation = Presentation(BytesIO(content))
    slides = []
    for slide in presentation.slides:
        text = []
        for shape in slide.shapes:
            value = getattr(shape, "text", None)
            if isinstance(value, str) and value.strip():
                text.append(value)
        if text:
            slides.append("\n".join(text))
    return "\n\n".join(slides).strip()
