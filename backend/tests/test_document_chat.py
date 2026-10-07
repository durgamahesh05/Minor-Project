from types import SimpleNamespace
from io import BytesIO
from unittest.mock import Mock, patch

import pytest
from bson import ObjectId
from fastapi import HTTPException, UploadFile

with patch("pymongo.MongoClient"):
    import main
from rag.pipeline import MultimodalRAG


@pytest.mark.parametrize("status", ["uploading", "processing", "failed", "ready"])
def test_chat_rejects_unindexed_document(monkeypatch, status):
    monkeypatch.setattr(main, "owned", lambda collection, ident, user: {"_id": ObjectId(), "status": status})
    rag = Mock()
    monkeypatch.setattr(main, "get_rag", rag)
    with pytest.raises(HTTPException) as error:
        main.message("conversation", main.Data(text="Explain this file", documentId=str(ObjectId())), {"_id": ObjectId()})
    assert error.value.status_code == 409
    rag.assert_not_called()


def test_chat_blocks_unscoped_request_during_indexing(monkeypatch):
    monkeypatch.setattr(main, "owned", lambda *args: {"_id": ObjectId()})
    database = Mock()
    database.documents.find_one.return_value = {"status": "processing"}
    monkeypatch.setattr(main, "db", database)
    with pytest.raises(HTTPException) as error:
        main.message("conversation", main.Data(text="Explain this file"), {"_id": ObjectId()})
    assert error.value.status_code == 409


def test_indexed_document_is_passed_to_rag(monkeypatch):
    document_id, user_id = ObjectId(), ObjectId()
    monkeypatch.setattr(main, "owned", lambda *args: {"_id": document_id, "status": "indexed"})
    database = Mock()
    database.messages.insert_one.return_value.inserted_id = ObjectId()
    monkeypatch.setattr(main, "db", database)
    rag = Mock()
    rag.chat.return_value.answer = "Explanation"
    monkeypatch.setattr(main, "get_rag", lambda _: rag)
    main.message("conversation", main.Data(text="Explain this file", documentId=str(document_id)), {"_id": user_id})
    rag.chat.assert_called_once_with("Explain this file", user_id, document_id=document_id, response_level="simple")


def test_selected_document_bypasses_general_router():
    rag = MultimodalRAG.__new__(MultimodalRAG)
    rag.answer, rag.route = Mock(), Mock()
    user_id, document_id = ObjectId(), ObjectId()
    rag.chat("read this file and explain it", user_id, document_id=document_id)
    rag.route.assert_not_called()
    rag.answer.assert_called_once_with("read this file and explain it", user_id, document_id=document_id, response_level="simple")


def test_selected_document_retrieval_keeps_low_similarity_and_scopes_both_stores():
    rag = MultimodalRAG.__new__(MultimodalRAG)
    rag.vectorstore, rag.embedder, rag.database = Mock(), Mock(), Mock()
    user_id, document_id = ObjectId(), ObjectId()
    rag.embedder.embed_query.return_value = [0.1]
    rag.vectorstore.search.return_value = [SimpleNamespace(score=0.01, payload={"element_id": "element"})]
    rag.database.rag_elements.find.return_value = [{"_id": "element", "content": "Document text"}]
    result = rag.retrieve("read this file", user_id, document_id=document_id)
    assert len(result) == 1
    rag.vectorstore.search.assert_called_once_with([0.1], str(user_id), 5, document_id=str(document_id))
    rag.database.rag_elements.find.assert_called_once_with({"_id": {"$in": ["element"]}, "userId": user_id, "documentId": document_id})


def test_message_saves_safe_attachment_snapshot(monkeypatch):
    uid, did, cid = ObjectId(), ObjectId(), ObjectId()
    document = {"_id": did, "status": "indexed", "originalName": "notes.pdf", "size": 123,
                "mimeType": "application/pdf", "clientDocumentId": "local-id",
                "storageProvider": "browser-opfs", "content": "must not be copied"}
    monkeypatch.setattr(main, "owned", lambda collection, *args: document if collection == "documents" else {"_id": cid})
    db, rag = Mock(), Mock()
    db.messages.insert_one.return_value.inserted_id = ObjectId()
    rag.chat.return_value.answer = "Here is the summary."
    monkeypatch.setattr(main, "db", db)
    monkeypatch.setattr(main, "get_rag", lambda _: rag)
    result = main.message(str(cid), main.Data(text="Read this", documentId=str(did)), {"_id": uid})
    attachment = result["userMessage"]["attachment"]
    assert attachment["id"] == str(did)
    assert attachment["originalName"] == "notes.pdf"
    assert "content" not in attachment


def test_upload_marks_indexed_only_after_ingestion_returns(monkeypatch):
    database = Mock()
    database.documents.insert_one.return_value.inserted_id = ObjectId()
    monkeypatch.setattr(main, "db", database)
    rag = Mock()

    def ingest(*args):
        database.documents.update_one.assert_not_called()
        return 3

    rag.ingest.side_effect = ingest
    monkeypatch.setattr(main, "get_rag", lambda _: rag)
    result = main.upload(UploadFile(filename="notes.txt", file=BytesIO(b"Some notes")), "12345678-1234-1234-1234-123456789012", {"_id": ObjectId()})
    assert result["document"]["status"] == "indexed"
    assert result["document"]["elementCount"] == 3
    assert database.documents.insert_one.call_args.args[0]["storageProvider"] == "browser-opfs"
    assert any(call.args[1]["$set"].get("status") == "indexed" for call in database.documents.update_one.call_args_list)


def test_failed_ingestion_never_marks_upload_indexed(monkeypatch):
    database = Mock()
    database.documents.insert_one.return_value.inserted_id = ObjectId()
    monkeypatch.setattr(main, "db", database)
    rag = Mock()
    rag.ingest.side_effect = RuntimeError("Qdrant unavailable")
    monkeypatch.setattr(main, "get_rag", lambda _: rag)
    cleanup = Mock()
    monkeypatch.setattr(main, "delete_document_artifacts", cleanup)
    with pytest.raises(HTTPException) as error:
        main.upload(UploadFile(filename="notes.txt", file=BytesIO(b"Some notes")), "12345678-1234-1234-1234-123456789012", {"_id": ObjectId()})
    assert error.value.status_code == 502
    database.documents.update_one.assert_not_called()
    cleanup.assert_called_once()


def test_docx_content_reaches_extraction():
    from docx import Document
    from ingestion.file_router import extract_text

    document = Document()
    document.add_paragraph("Synapse requirements")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Users can upload study notes"
    content = BytesIO()
    document.save(content)
    result = extract_text("requirements.docx", content.getvalue())
    assert "Synapse requirements" in result
    assert "Users can upload study notes" in result


def test_text_ingestion_indexes_chunks_without_network_summarization():
    rag = MultimodalRAG.__new__(MultimodalRAG)
    rag.summary_chain, rag.embedder, rag.vectorstore, rag.database = Mock(), Mock(), Mock(), Mock()
    rag.embedder.embed.return_value = [[0.1, 0.2]]
    count = rag.ingest("notes.txt", b"Specific requirements and facts.", ObjectId(), ObjectId())
    assert count == 1
    rag.summary_chain.invoke.assert_not_called()
    rag.embedder.embed.assert_called_once_with(["Specific requirements and facts."])
    assert rag.vectorstore.upsert.call_args.args[0][0]["payload"]["summary"] == "Specific requirements and facts."


@pytest.mark.parametrize("title", ["", "   ", "x" * 81])
def test_rename_rejects_invalid_titles(monkeypatch, title):
    database = Mock()
    monkeypatch.setattr(main, "db", database)
    with pytest.raises(HTTPException) as error:
        main.rename_chat(str(ObjectId()), main.Data(title=title), {"_id": ObjectId()})
    assert error.value.status_code == 400
    database.conversations.update_one.assert_not_called()


def test_rename_persists_trimmed_title_for_owner(monkeypatch):
    database = Mock()
    user_id, conversation_id = ObjectId(), ObjectId()
    database.__getitem__ = Mock(return_value=database.conversations)
    database.conversations.find_one.return_value = {"_id": conversation_id, "userId": user_id, "title": "Old name"}
    monkeypatch.setattr(main, "db", database)
    result = main.rename_chat(str(conversation_id), main.Data(title="  Study notes  "), {"_id": user_id})
    assert result["conversation"]["title"] == "Study notes"
    database.conversations.update_one.assert_called_once_with(
        {"_id": conversation_id, "userId": user_id}, {"$set": {"title": "Study notes"}}
    )


def test_rename_cannot_modify_another_users_chat(monkeypatch):
    database = Mock()
    database.__getitem__ = Mock(return_value=database.conversations)
    database.conversations.find_one.return_value = None
    monkeypatch.setattr(main, "db", database)
    with pytest.raises(HTTPException) as error:
        main.rename_chat(str(ObjectId()), main.Data(title="Changed"), {"_id": ObjectId()})
    assert error.value.status_code == 404
    database.conversations.update_one.assert_not_called()
