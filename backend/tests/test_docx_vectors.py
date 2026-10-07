from io import BytesIO
from unittest.mock import Mock

from bson import ObjectId
from docx import Document
from PIL import Image

from ingestion.docx_reader import read_docx
from ingestion.multimodal import extract_elements
from rag.pipeline import MultimodalRAG
from vectorstore.chroma import ChromaVectorStore


def word_file(image_only=False):
    document = Document()
    if not image_only:
        document.add_paragraph("Before the table")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Term"
        table.cell(0, 1).text = "Definition"
        document.add_paragraph("After the table")
        document.sections[0].header.paragraphs[0].text = "Study guide"
    image = BytesIO()
    Image.new("RGB", (40, 40), "white").save(image, "PNG")
    document.add_picture(BytesIO(image.getvalue()))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_preserves_table_order_and_headers():
    text = read_docx(word_file())
    assert text.index("Before the table") < text.index("Definition") < text.index("After the table")
    assert "Study guide" in text


def test_image_only_docx_reaches_vision(monkeypatch):
    monkeypatch.setattr("ingestion.multimodal.cloud_ocr_configured", lambda: False)
    elements = extract_elements("scan.docx", word_file(image_only=True))
    assert len(elements) == 1
    assert elements[0].kind == "image"


def test_docx_ingestion_stores_vectors_in_real_chroma(tmp_path, monkeypatch):
    monkeypatch.setenv("CHROMA_MODE", "local")
    monkeypatch.setenv("ENVIRONMENT", "development")
    monkeypatch.setenv("CHROMA_PATH", str(tmp_path / "chroma"))
    monkeypatch.setattr("ingestion.multimodal.cloud_ocr_configured", lambda: False)
    rag = MultimodalRAG.__new__(MultimodalRAG)
    rag.database = Mock()
    # Deterministic vectors isolate persistence from model downloads/API calls.
    rag.embedder = Mock()
    rag.embedder.embed.side_effect = lambda texts: [[1.0, 0.0] for _ in texts]
    rag._summary = lambda element: "Diagram of a cell"
    rag.vectorstore = ChromaVectorStore(2)
    uid, did = ObjectId(), ObjectId()
    count = rag.ingest("notes.docx", word_file(), uid, did)
    reopened = ChromaVectorStore(2)
    rows = reopened.collection.get(where={"document_id": str(did)}, include=["embeddings", "metadatas", "documents"])
    assert len(rows["ids"]) == count >= 2
    assert all(metadata["user_id"] == str(uid) for metadata in rows["metadatas"])
    assert any("Definition" in text for text in rows["documents"])
    assert all(list(vector) == [1.0, 0.0] for vector in rows["embeddings"])
    assert reopened.search([1.0, 0.0], str(uid), document_id=str(did))
